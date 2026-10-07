"""OCR for image-only slides / scanned documents."""
import io

import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

from uni_agent.extract import OCR_MARK, extract_text, has_real_text
from uni_agent.ocr import get_ocr


def png(text: str = "", size=(1200, 300)) -> bytes:
    img = Image.new("RGB", size, "white")
    if text:
        ImageDraw.Draw(img).text((40, 100), text, fill="black", font=ImageFont.load_default(size=56))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


class FakeOCR:
    def __init__(self, text="Normalization removes redundancy and update anomalies from relations. " * 2):
        self.text, self.calls = text, 0

    def __call__(self, data: bytes) -> str:
        self.calls += 1
        return self.text


def make_pptx(tmp_path, slides):
    """slides: list of ('image', None) or ('text', 'some text')."""
    prs = Presentation()
    for kind, text in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        if kind == "image":
            slide.shapes.add_picture(io.BytesIO(png()), Inches(0), Inches(0), width=Inches(10))
        else:
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.text = text
            slide.shapes.add_picture(io.BytesIO(png()), Inches(1), Inches(4), width=Inches(2))  # a logo
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return path


def test_image_only_slides_are_ocred_and_marked(tmp_path):
    path = make_pptx(tmp_path, [("image", None), ("image", None)])
    assert not has_real_text(extract_text(path))  # without OCR: nothing to summarise
    ocr = FakeOCR()
    text = extract_text(path, ocr)
    assert ocr.calls == 2 and text.count(OCR_MARK) == 2
    assert "Normalization removes redundancy" in text and has_real_text(text)


def test_slides_with_real_text_are_not_ocred(tmp_path):
    path = make_pptx(tmp_path, [("text", "Functional dependencies: X determines Y when equal X values imply equal Y.")])
    ocr = FakeOCR()
    text = extract_text(path, ocr)
    assert ocr.calls == 0 and OCR_MARK not in text  # the logo picture is ignored


def test_mixed_deck_only_ocrs_image_slides(tmp_path):
    path = make_pptx(tmp_path, [("text", "A slide with plenty of genuine text on it for sure."), ("image", None)])
    ocr = FakeOCR()
    text = extract_text(path, ocr)
    assert ocr.calls == 1
    slide1, slide2 = text.split("## Slide 2")
    assert OCR_MARK not in slide1 and OCR_MARK in slide2


def test_blank_ocr_result_adds_nothing(tmp_path):
    path = make_pptx(tmp_path, [("image", None)])
    assert OCR_MARK not in extract_text(path, FakeOCR(text="   "))


def test_docx_of_screenshots_is_ocred(tmp_path):
    doc = Document()
    doc.add_picture(io.BytesIO(png()))
    path = tmp_path / "scan.docx"
    doc.save(path)
    ocr = FakeOCR()
    assert OCR_MARK in extract_text(path, ocr) and ocr.calls == 1


def test_docx_with_text_is_not_ocred(tmp_path):
    doc = Document()
    doc.add_paragraph("Lab 3 asks you to create the tables and insert sample rows. " * 5)
    doc.add_picture(io.BytesIO(png()))
    path = tmp_path / "sheet.docx"
    doc.save(path)
    ocr = FakeOCR()
    extract_text(path, ocr)
    assert ocr.calls == 0


@pytest.mark.skipif(get_ocr() is None, reason="Windows OCR not available")
def test_windows_ocr_reads_rendered_text(tmp_path):
    ocr = get_ocr()
    assert "normalization" in ocr(png("Database Normalization")).lower()
    assert ocr(png()) == ""  # blank image
    assert ocr(b"not an image") == ""  # never raises
