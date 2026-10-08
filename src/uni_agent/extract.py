"""File -> plain text (pdf/pptx/docx/txt/md), file hashing, and section-aware chunking."""
from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Callable

SUPPORTED = {".pdf", ".pptx", ".docx", ".txt", ".md"}

OcrFunc = Callable[[bytes], str]
OCR_MARK = "[OCR from image]"
VOICE_MARK = "[Voice transcript]"
OCR_BELOW_CHARS = 40  # a slide/page with less real text than this is treated as an image


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _ocr_block(blobs: list[bytes], ocr: OcrFunc) -> str:
    texts = [t.strip() for t in (ocr(b) for b in blobs) if t.strip()]
    return f"{OCR_MARK}\n" + "\n".join(texts) if texts else ""


def _pdf(path: Path, ocr: OcrFunc | None) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if ocr and len(text) < OCR_BELOW_CHARS:  # scanned page: the text is in the page image(s)
            try:
                blobs = [img.data for img in page.images]
            except Exception:  # image encodings pypdf can't decode
                blobs = []
            text = "\n".join(x for x in (text, _ocr_block(blobs, ocr)) if x)
        parts.append(f"## Page {i}\n{text}")
    return "\n\n".join(parts)


def _picture_blob(shape) -> bytes | None:
    try:
        return shape.image.blob
    except (AttributeError, ValueError, KeyError):  # empty placeholder / linked image
        return None


def _pptx(path: Path, ocr: OcrFunc | None) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    parts = []
    for i, slide in enumerate(Presentation(str(path)).slides, 1):
        lines = []
        pictures = []
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):  # incl. picture placeholders
                pictures.append(shape)
            elif shape.has_text_frame:
                lines += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        if ocr and pictures and len("".join(lines)) < OCR_BELOW_CHARS:  # slide is a screenshot/scan
            block = _ocr_block([b for b in map(_picture_blob, pictures) if b], ocr)
            if block:
                lines.append(block)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[Speaker notes] {notes}")
        parts.append(f"## Slide {i}\n" + "\n".join(lines))
    return "\n\n".join(parts)


def _docx(path: Path, ocr: OcrFunc | None) -> str:
    from docx import Document

    doc = Document(str(path))
    lines = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        lines.append(f"## {p.text}" if style.startswith("heading") or style == "title" else p.text)
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    if ocr and len("".join(lines)) < OCR_BELOW_CHARS * 5:  # mostly pasted screenshots/scans
        blobs = [rel.target_part.blob for rel in doc.part.rels.values() if "image" in rel.reltype]
        block = _ocr_block(blobs, ocr)
        if block:
            lines.append(block)
    return "\n".join(lines)


def extract_text(path: Path, ocr: OcrFunc | None = None) -> str:
    """With `ocr`, slides/pages/documents that are mostly images also get their image text read (marked)."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf(path, ocr)
    if ext == ".pptx":
        return _pptx(path, ocr)
    if ext == ".docx":
        return _docx(path, ocr)
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {ext}")


def has_real_text(text: str, min_chars: int = 200) -> bool:
    """False for scanned PDFs / image-only slides where extraction yields only headers."""
    body = "\n".join(l for l in text.splitlines() if not l.startswith("## "))
    return len(body.strip()) >= min_chars


def chunk_text(text: str, max_chars: int) -> list[str]:
    """Split on '## ' section headers (slides/pages/headings), packing sections up to max_chars.
    A single oversized section is split on line boundaries."""
    sections: list[str] = []
    current: list[str] = []
    for line in text.splitlines():
        if line.startswith("## ") and current:
            sections.append("\n".join(current))
            current = []
        current.append(line)
    if current:
        sections.append("\n".join(current))

    pieces: list[str] = []
    for sec in sections:
        while len(sec) > max_chars:
            cut = sec.rfind("\n", 0, max_chars)
            cut = cut if cut > 0 else max_chars
            pieces.append(sec[:cut])
            sec = sec[cut:].lstrip("\n")
        pieces.append(sec)

    chunks: list[str] = []
    buf = ""
    for piece in pieces:
        if buf and len(buf) + len(piece) + 2 > max_chars:
            chunks.append(buf)
            buf = piece
        else:
            buf = f"{buf}\n\n{piece}" if buf else piece
    if buf.strip():
        chunks.append(buf)
    return chunks
